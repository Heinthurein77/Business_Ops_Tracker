"""ERP Phase 1 — accounting core: chart of accounts, double-entry journal,
trial balance, ledger, P&L, balance sheet, fiscal periods, posting rules,
opening balances.

Rules enforced here:
- Every posted entry balances (Σ debit = Σ credit), each line one-sided.
- Posted entries are immutable; corrections use reversal entries.
- Posting into a closed fiscal period is rejected.
- Auto-posting (daily financials, approved expenses) goes through posting
  rules: BU rule → group rule → built-in defaults from accounting_templates.
"""
from datetime import date, datetime
from decimal import Decimal
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlalchemy.orm import Session

from src.accounting_templates import (
    DEFAULT_POSTING_RULES, EXPENSE_CATEGORY_ACCOUNTS, TEMPLATE_TYPES,
    template_accounts,
)
from src.auth import get_current_user, require_roles
from src.database import get_db
from src.models import (
    Account, BusinessUnit, FiscalPeriod, JournalEntry, JournalLine,
    PostingRule, User,
)

router = APIRouter(prefix="/api/accounting", tags=["accounting"])
require_accounting = require_roles("Chairman", "HO", "Accountant")
require_admin = require_roles("Chairman", "HO")

ZERO = Decimal("0")


# ---------------------------------------------------------------- helpers

def period_of(d: date) -> str:
    return d.strftime("%Y-%m")


def fmt(v) -> float:
    return round(float(v or 0), 2)


def bu_or_404(db: Session, bu_id: int) -> BusinessUnit:
    bu = db.query(BusinessUnit).filter(BusinessUnit.id == bu_id).first()
    if not bu:
        raise HTTPException(404, "Business unit not found")
    return bu


def check_bu_access(user: User, bu_id: int) -> None:
    """CEOs only ever touch their own BU; accounting roles see all BUs."""
    if user.role == "CEO" and user.business_unit_id != bu_id:
        raise HTTPException(403, "Not your business unit")


def ensure_period_open(db: Session, bu_id: int, d: date) -> None:
    p = period_of(d)
    row = db.query(FiscalPeriod).filter(
        FiscalPeriod.business_unit_id == bu_id, FiscalPeriod.period == p).first()
    if row and row.status == "closed":
        raise HTTPException(400, f"Fiscal period {p} is closed for this business unit")


def account_by_code(db: Session, bu_id: int, code: str) -> Optional[Account]:
    return db.query(Account).filter(
        Account.business_unit_id == bu_id, Account.code == code,
        Account.is_active == True).first()  # noqa: E712


def next_entry_no(db: Session, bu_id: int, d: date) -> str:
    p = period_of(d)
    count = db.query(JournalEntry).filter(
        JournalEntry.business_unit_id == bu_id,
        JournalEntry.entry_no.like(f"JV-{p}-%")).count()
    return f"JV-{p}-{count + 1:03d}"


def serialize_account(a: Account) -> dict:
    return {
        "id": a.id, "business_unit_id": a.business_unit_id, "code": a.code,
        "name": a.name, "account_type": a.account_type, "parent_id": a.parent_id,
        "is_header": a.is_header, "is_system": a.is_system, "is_active": a.is_active,
    }


def serialize_entry(db: Session, e: JournalEntry, with_lines: bool = True) -> dict:
    out = {
        "id": e.id, "business_unit_id": e.business_unit_id,
        "entry_no": e.entry_no,
        "entry_date": e.entry_date.isoformat() if e.entry_date else None,
        "description": e.description, "reference": e.reference,
        "source": e.source, "source_id": e.source_id, "status": e.status,
        "reversal_of_id": e.reversal_of_id, "reversed_by_id": e.reversed_by_id,
        "total_debit": fmt(sum((l.debit or ZERO) for l in e.lines)),
        "total_credit": fmt(sum((l.credit or ZERO) for l in e.lines)),
        "created_at": e.created_at.isoformat() if e.created_at else None,
        "posted_at": e.posted_at.isoformat() if e.posted_at else None,
    }
    if with_lines:
        out["lines"] = [{
            "id": l.id, "account_id": l.account_id,
            "account_code": l.account.code if l.account else None,
            "account_name": l.account.name if l.account else None,
            "debit": fmt(l.debit), "credit": fmt(l.credit), "memo": l.memo,
        } for l in e.lines]
    return out


def validate_lines(db: Session, bu_id: int, lines: list[dict]) -> None:
    if len(lines) < 2:
        raise HTTPException(400, "A journal entry needs at least two lines")
    total_dr = ZERO
    total_cr = ZERO
    for ln in lines:
        dr = Decimal(str(ln.get("debit") or 0))
        cr = Decimal(str(ln.get("credit") or 0))
        if dr < 0 or cr < 0:
            raise HTTPException(400, "Amounts cannot be negative")
        if (dr > 0) == (cr > 0):
            raise HTTPException(400, "Each line must be debit-only or credit-only")
        total_dr += dr
        total_cr += cr
        acc = db.query(Account).filter(Account.id == ln["account_id"]).first()
        if not acc or acc.business_unit_id != bu_id:
            raise HTTPException(400, "Line account does not belong to this business unit")
        if acc.is_header:
            raise HTTPException(400, f"Cannot post to header account {acc.code}")
        if not acc.is_active:
            raise HTTPException(400, f"Account {acc.code} is inactive")
    if total_dr != total_cr:
        raise HTTPException(400, f"Entry does not balance: debit {total_dr} != credit {total_cr}")
    if total_dr == 0:
        raise HTTPException(400, "Entry total cannot be zero")


def create_entry(db: Session, bu_id: int, entry_date: date, description: str,
                 source: str, user_id: Optional[int], lines: list[dict],
                 reference: Optional[str] = None, source_id: Optional[int] = None,
                 post: bool = False, reversal_of_id: Optional[int] = None) -> JournalEntry:
    """Create an entry (draft by default). Caller commits."""
    ensure_period_open(db, bu_id, entry_date)
    if post:
        validate_lines(db, bu_id, lines)
    e = JournalEntry(
        business_unit_id=bu_id, entry_no=next_entry_no(db, bu_id, entry_date),
        entry_date=entry_date, description=description, reference=reference,
        source=source, source_id=source_id,
        status="posted" if post else "draft",
        reversal_of_id=reversal_of_id, created_by_id=user_id,
        posted_by_id=user_id if post else None,
        posted_at=datetime.utcnow() if post else None,
    )
    db.add(e)
    db.flush()
    for ln in lines:
        db.add(JournalLine(
            entry_id=e.id, account_id=ln["account_id"],
            debit=Decimal(str(ln.get("debit") or 0)),
            credit=Decimal(str(ln.get("credit") or 0)),
            memo=ln.get("memo"),
        ))
    db.flush()
    return e


# ------------------------------------------------------- auto-posting service

def resolve_rule(db: Session, bu_id: int, source_field: str) -> Optional[tuple[str, str]]:
    rule = db.query(PostingRule).filter(
        PostingRule.business_unit_id == bu_id,
        PostingRule.source_field == source_field).first()
    if rule:
        return rule.debit_account_code, rule.credit_account_code
    rule = db.query(PostingRule).filter(
        PostingRule.business_unit_id == None,  # noqa: E711
        PostingRule.source_field == source_field).first()
    if rule:
        return rule.debit_account_code, rule.credit_account_code
    for field, dr, cr in DEFAULT_POSTING_RULES:
        if field == source_field:
            return dr, cr
    return None


def _rule_lines(db: Session, bu_id: int, source_field: str, amount: Decimal,
                debit_override: Optional[str] = None):
    codes = resolve_rule(db, bu_id, source_field)
    if not codes:
        return None
    dr_code, cr_code = codes
    if debit_override:
        dr_code = debit_override
    dr_acc = account_by_code(db, bu_id, dr_code)
    cr_acc = account_by_code(db, bu_id, cr_code)
    if not dr_acc or not cr_acc:
        return None  # BU not provisioned yet — auto-posting silently skips
    return [
        {"account_id": dr_acc.id, "debit": amount, "credit": 0},
        {"account_id": cr_acc.id, "debit": 0, "credit": amount},
    ]


def auto_post_daily_financial(db: Session, df, user_id: Optional[int]) -> None:
    """Post journal entries for a DailyFinancial row (revenue + expenses).

    Cash flow / receivable / payable daily figures are position snapshots,
    not transactions — they are NOT posted here (AR/AP modules, Phase 4).
    Re-submitting a day reverses the previous auto entries and reposts.
    """
    bu_id = df.business_unit_id
    if not db.query(Account).filter(Account.business_unit_id == bu_id).first():
        return  # BU has no CoA yet
    old = db.query(JournalEntry).filter(
        JournalEntry.source == "daily_financial", JournalEntry.source_id == df.id,
        JournalEntry.status == "posted", JournalEntry.reversed_by_id == None).all()  # noqa: E711
    for e in old:
        _reverse_entry(db, e, user_id)
    for field, amount in (("daily_financial.revenue", df.revenue),
                          ("daily_financial.expenses", df.expenses)):
        amt = Decimal(str(amount or 0))
        if amt <= 0:
            continue
        lines = _rule_lines(db, bu_id, field, amt)
        if not lines:
            continue
        create_entry(db, bu_id, df.report_date,
                     f"Daily report {df.report_date.isoformat()} — {field.split('.')[-1]}",
                     "daily_financial", user_id, lines, source_id=df.id, post=True)
    db.commit()


def auto_post_expense(db: Session, expense, user_id: Optional[int]) -> None:
    """Post an approved ExpenseRequest: Dr expense account (by category), Cr cash."""
    bu_id = expense.business_unit_id
    amt = Decimal(str(expense.amount or 0))
    if amt <= 0:
        return
    override = None
    if expense.category:
        override = EXPENSE_CATEGORY_ACCOUNTS.get(expense.category.strip().lower())
    lines = _rule_lines(db, bu_id, "expense_request", amt, debit_override=override)
    if not lines:
        return
    create_entry(db, bu_id, expense.created_at.date() if expense.created_at else date.today(),
                 f"Expense approved: {expense.title}", "expense", user_id, lines,
                 source_id=expense.id, post=True)
    db.commit()


def _reverse_entry(db: Session, entry: JournalEntry, user_id: Optional[int]) -> JournalEntry:
    lines = [{"account_id": l.account_id, "debit": l.credit, "credit": l.debit,
              "memo": l.memo} for l in entry.lines]
    rev = create_entry(db, entry.business_unit_id, date.today(),
                       f"Reversal of {entry.entry_no}", "reversal", user_id,
                       lines, post=True, reversal_of_id=entry.id)
    entry.reversed_by_id = rev.id
    return rev


# ---------------------------------------------------------------- templates & CoA

@router.get("/templates")
def list_templates(user: User = Depends(require_accounting)):
    return {"templates": TEMPLATE_TYPES}


@router.post("/provision/{bu_id}")
def provision_bu(bu_id: int, template: str = "retail",
                 user: User = Depends(require_accounting), db: Session = Depends(get_db)):
    """Clone a CoA template into a BU + seed group posting rules. Idempotent."""
    bu_or_404(db, bu_id)
    if template not in TEMPLATE_TYPES:
        raise HTTPException(400, f"template must be one of {TEMPLATE_TYPES}")
    existing = db.query(Account).filter(Account.business_unit_id == bu_id).count()
    if existing:
        return {"ok": True, "created": 0, "message": "BU already has a chart of accounts"}
    id_by_code: dict[str, int] = {}
    created = 0
    for code, name, acc_type, parent_code, is_header in template_accounts(template):
        acc = Account(business_unit_id=bu_id, code=code, name=name,
                      account_type=acc_type, is_header=is_header, is_system=True,
                      parent_id=id_by_code.get(parent_code) if parent_code else None)
        db.add(acc)
        db.flush()
        id_by_code[code] = acc.id
        created += 1
    for field, dr, cr in DEFAULT_POSTING_RULES:
        if not db.query(PostingRule).filter(
                PostingRule.business_unit_id == None,  # noqa: E711
                PostingRule.source_field == field).first():
            db.add(PostingRule(business_unit_id=None, source_field=field,
                               debit_account_code=dr, credit_account_code=cr))
    db.commit()
    return {"ok": True, "created": created}


@router.get("/coa")
def get_coa(business_unit_id: int, user: User = Depends(get_current_user),
            db: Session = Depends(get_db)):
    check_bu_access(user, business_unit_id)
    rows = db.query(Account).filter(
        Account.business_unit_id == business_unit_id).order_by(Account.code).all()
    return {"items": [serialize_account(a) for a in rows]}


class AccountIn(BaseModel):
    business_unit_id: int
    code: str = Field(min_length=1, max_length=20)
    name: str = Field(min_length=1, max_length=150)
    account_type: str
    parent_id: Optional[int] = None
    is_header: bool = False


@router.post("/accounts")
def create_account(payload: AccountIn, user: User = Depends(require_accounting),
                   db: Session = Depends(get_db)):
    bu_or_404(db, payload.business_unit_id)
    if payload.account_type not in ("asset", "liability", "equity", "income", "expense"):
        raise HTTPException(400, "Invalid account_type")
    if account_by_code(db, payload.business_unit_id, payload.code):
        raise HTTPException(400, f"Account code {payload.code} already exists in this BU")
    acc = Account(business_unit_id=payload.business_unit_id, code=payload.code.strip(),
                  name=payload.name, account_type=payload.account_type,
                  parent_id=payload.parent_id, is_header=payload.is_header)
    db.add(acc)
    db.commit()
    db.refresh(acc)
    return serialize_account(acc)


class AccountUpdate(BaseModel):
    name: Optional[str] = None
    is_active: Optional[bool] = None


@router.patch("/accounts/{account_id}")
def update_account(account_id: int, payload: AccountUpdate,
                   user: User = Depends(require_accounting), db: Session = Depends(get_db)):
    acc = db.query(Account).filter(Account.id == account_id).first()
    if not acc:
        raise HTTPException(404, "Account not found")
    if payload.name is not None:
        acc.name = payload.name
    if payload.is_active is not None:
        acc.is_active = payload.is_active
    db.commit()
    db.refresh(acc)
    return serialize_account(acc)


# ---------------------------------------------------------------- journal

class JournalLineIn(BaseModel):
    account_id: int
    debit: float = 0
    credit: float = 0
    memo: Optional[str] = None


class JournalIn(BaseModel):
    business_unit_id: int
    entry_date: date
    description: Optional[str] = None
    reference: Optional[str] = None
    lines: list[JournalLineIn]


@router.post("/journal")
def create_journal(payload: JournalIn, user: User = Depends(require_accounting),
                   db: Session = Depends(get_db)):
    bu_or_404(db, payload.business_unit_id)
    lines = [ln.model_dump() for ln in payload.lines]
    e = create_entry(db, payload.business_unit_id, payload.entry_date,
                     payload.description or "", "manual", user.id, lines,
                     reference=payload.reference)
    db.commit()
    db.refresh(e)
    return serialize_entry(db, e)


@router.get("/journal")
def list_journal(business_unit_id: int, period: Optional[str] = None,
                 status: Optional[str] = None, source: Optional[str] = None,
                 user: User = Depends(require_accounting), db: Session = Depends(get_db)):
    q = db.query(JournalEntry).filter(JournalEntry.business_unit_id == business_unit_id)
    if status:
        q = q.filter(JournalEntry.status == status)
    if source:
        q = q.filter(JournalEntry.source == source)
    rows = q.order_by(JournalEntry.entry_date.desc(), JournalEntry.id.desc()).limit(300).all()
    if period:
        rows = [r for r in rows if period_of(r.entry_date) == period]
    return {"items": [serialize_entry(db, r, with_lines=False) for r in rows]}


@router.get("/journal/{entry_id}")
def get_journal(entry_id: int, user: User = Depends(require_accounting),
                db: Session = Depends(get_db)):
    e = db.query(JournalEntry).filter(JournalEntry.id == entry_id).first()
    if not e:
        raise HTTPException(404, "Journal entry not found")
    return serialize_entry(db, e)


@router.post("/journal/{entry_id}/post")
def post_journal(entry_id: int, user: User = Depends(require_accounting),
                 db: Session = Depends(get_db)):
    e = db.query(JournalEntry).filter(JournalEntry.id == entry_id).first()
    if not e:
        raise HTTPException(404, "Journal entry not found")
    if e.status == "posted":
        raise HTTPException(400, "Entry already posted")
    ensure_period_open(db, e.business_unit_id, e.entry_date)
    lines = [{"account_id": l.account_id, "debit": l.debit, "credit": l.credit}
             for l in e.lines]
    validate_lines(db, e.business_unit_id, lines)
    e.status = "posted"
    e.posted_by_id = user.id
    e.posted_at = datetime.utcnow()
    db.commit()
    db.refresh(e)
    return serialize_entry(db, e)


@router.post("/journal/{entry_id}/reverse")
def reverse_journal(entry_id: int, user: User = Depends(require_accounting),
                    db: Session = Depends(get_db)):
    e = db.query(JournalEntry).filter(JournalEntry.id == entry_id).first()
    if not e:
        raise HTTPException(404, "Journal entry not found")
    if e.status != "posted":
        raise HTTPException(400, "Only posted entries can be reversed")
    if e.reversed_by_id:
        raise HTTPException(400, "Entry already reversed")
    rev = _reverse_entry(db, e, user.id)
    db.commit()
    db.refresh(rev)
    return serialize_entry(db, rev)


@router.delete("/journal/{entry_id}")
def delete_journal(entry_id: int, user: User = Depends(require_accounting),
                   db: Session = Depends(get_db)):
    e = db.query(JournalEntry).filter(JournalEntry.id == entry_id).first()
    if not e:
        raise HTTPException(404, "Journal entry not found")
    if e.status == "posted":
        raise HTTPException(400, "Posted entries cannot be deleted — reverse instead")
    db.delete(e)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- reports

def _posted_lines(db: Session, bu_id: int, upto: Optional[date] = None):
    q = (db.query(JournalLine, Account, JournalEntry)
         .join(Account, JournalLine.account_id == Account.id)
         .join(JournalEntry, JournalLine.entry_id == JournalEntry.id)
         .filter(JournalEntry.business_unit_id == bu_id,
                 JournalEntry.status == "posted"))
    if upto:
        q = q.filter(JournalEntry.entry_date <= upto)
    return q.all()


def _period_end(period: Optional[str]) -> Optional[date]:
    if not period:
        return None
    y, m = period.split("-")
    import calendar
    return date(int(y), int(m), calendar.monthrange(int(y), int(m))[1])


@router.get("/trial-balance")
def trial_balance(business_unit_id: int, period: Optional[str] = None,
                  user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    check_bu_access(user, business_unit_id)
    totals: dict[int, dict] = {}
    for line, acc, _e in _posted_lines(db, business_unit_id, _period_end(period)):
        t = totals.setdefault(acc.id, {"account": acc, "debit": ZERO, "credit": ZERO})
        t["debit"] += line.debit or ZERO
        t["credit"] += line.credit or ZERO
    items = []
    td = tc = ZERO
    for t in totals.values():
        bal = t["debit"] - t["credit"]
        items.append({
            "account_id": t["account"].id, "code": t["account"].code,
            "name": t["account"].name, "account_type": t["account"].account_type,
            "debit": fmt(t["debit"]), "credit": fmt(t["credit"]),
            "balance": fmt(bal),
            "balance_side": "debit" if bal >= 0 else "credit",
        })
        td += t["debit"]
        tc += t["credit"]
    items.sort(key=lambda x: x["code"])
    return {"items": items, "total_debit": fmt(td), "total_credit": fmt(tc),
            "balanced": td == tc}


@router.get("/ledger")
def ledger(account_id: int, period: Optional[str] = None,
           user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    acc = db.query(Account).filter(Account.id == account_id).first()
    if not acc:
        raise HTTPException(404, "Account not found")
    check_bu_access(user, acc.business_unit_id)
    rows = (db.query(JournalLine, JournalEntry)
            .join(JournalEntry, JournalLine.entry_id == JournalEntry.id)
            .filter(JournalLine.account_id == account_id,
                    JournalEntry.status == "posted")
            .order_by(JournalEntry.entry_date, JournalEntry.id).all())
    if period:
        rows = [r for r in rows if period_of(r[1].entry_date) == period]
    balance = ZERO
    items = []
    for line, e in rows:
        balance += (line.debit or ZERO) - (line.credit or ZERO)
        items.append({
            "entry_id": e.id, "entry_no": e.entry_no,
            "date": e.entry_date.isoformat(), "description": e.description,
            "debit": fmt(line.debit), "credit": fmt(line.credit),
            "balance": fmt(balance), "memo": line.memo,
        })
    return {"account": serialize_account(acc), "items": items, "balance": fmt(balance)}


@router.get("/reports/pl")
def profit_loss(business_unit_id: int, period: Optional[str] = None,
                user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    check_bu_access(user, business_unit_id)
    income: dict[str, Decimal] = {}
    expense: dict[str, Decimal] = {}
    for line, acc, e in _posted_lines(db, business_unit_id, _period_end(period)):
        if period and period_of(e.entry_date) != period:
            continue
        if acc.account_type == "income":
            income[acc.code] = income.get(acc.code, ZERO) + (line.credit or ZERO) - (line.debit or ZERO)
        elif acc.account_type == "expense":
            expense[acc.code] = expense.get(acc.code, ZERO) + (line.debit or ZERO) - (line.credit or ZERO)
    names = {a.code: a.name for a in db.query(Account).filter(
        Account.business_unit_id == business_unit_id).all()}
    ti = sum(income.values(), ZERO)
    te = sum(expense.values(), ZERO)
    return {
        "income": [{"code": c, "name": names.get(c, c), "amount": fmt(v)} for c, v in sorted(income.items())],
        "expenses": [{"code": c, "name": names.get(c, c), "amount": fmt(v)} for c, v in sorted(expense.items())],
        "total_income": fmt(ti), "total_expenses": fmt(te), "net_profit": fmt(ti - te),
    }


@router.get("/reports/balance-sheet")
def balance_sheet(business_unit_id: int, period: Optional[str] = None,
                  user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    check_bu_access(user, business_unit_id)
    groups: dict[str, dict[str, Decimal]] = {"asset": {}, "liability": {}, "equity": {}}
    retained = ZERO
    for line, acc, _e in _posted_lines(db, business_unit_id, _period_end(period)):
        if acc.account_type in groups:
            bal = (line.debit or ZERO) - (line.credit or ZERO)
            if acc.account_type != "asset":
                bal = -bal  # liability/equity are credit-normal
            groups[acc.account_type][acc.code] = groups[acc.account_type].get(acc.code, ZERO) + bal
        elif acc.account_type == "income":
            retained += (line.credit or ZERO) - (line.debit or ZERO)
        elif acc.account_type == "expense":
            retained -= (line.debit or ZERO) - (line.credit or ZERO)
    names = {a.code: a.name for a in db.query(Account).filter(
        Account.business_unit_id == business_unit_id).all()}

    def section(kind):
        items = [{"code": c, "name": names.get(c, c), "amount": fmt(v)}
                 for c, v in sorted(groups[kind].items()) if v != 0]
        return items, sum(groups[kind].values(), ZERO)

    assets, ta = section("asset")
    liabs, tl = section("liability")
    equity, te = section("equity")
    te += retained
    equity.append({"code": "—", "name": "Retained Earnings (current result)",
                   "amount": fmt(retained)})
    return {
        "assets": assets, "total_assets": fmt(ta),
        "liabilities": liabs, "total_liabilities": fmt(tl),
        "equity": equity, "total_equity": fmt(te),
        "total_liabilities_equity": fmt(tl + te),
        "balanced": ta == tl + te,
    }


# ------------------------------------------------------- Phase 2: consolidation

def _all_posted_lines(db: Session, upto: Optional[date] = None):
    q = (db.query(JournalLine, Account, JournalEntry)
         .join(Account, JournalLine.account_id == Account.id)
         .join(JournalEntry, JournalLine.entry_id == JournalEntry.id)
         .filter(JournalEntry.status == "posted"))
    if upto:
        q = q.filter(JournalEntry.entry_date <= upto)
    return q.all()


@router.get("/reports/consolidated-pl")
def consolidated_pl(period: Optional[str] = None, user: User = Depends(require_accounting),
                    db: Session = Depends(get_db)):
    """Group P&L: every BU summed by the shared account-code standard."""
    income: dict[str, Decimal] = {}
    expense: dict[str, Decimal] = {}
    names: dict[str, str] = {}
    for line, acc, e in _all_posted_lines(db, _period_end(period)):
        names.setdefault(acc.code, acc.name)
        if period and period_of(e.entry_date) != period:
            continue
        if acc.account_type == "income":
            income[acc.code] = income.get(acc.code, ZERO) + (line.credit or ZERO) - (line.debit or ZERO)
        elif acc.account_type == "expense":
            expense[acc.code] = expense.get(acc.code, ZERO) + (line.debit or ZERO) - (line.credit or ZERO)
    ti = sum(income.values(), ZERO)
    te = sum(expense.values(), ZERO)
    return {
        "income": [{"code": c, "name": names.get(c, c), "amount": fmt(v)} for c, v in sorted(income.items())],
        "expenses": [{"code": c, "name": names.get(c, c), "amount": fmt(v)} for c, v in sorted(expense.items())],
        "total_income": fmt(ti), "total_expenses": fmt(te), "net_profit": fmt(ti - te),
    }


@router.get("/reports/consolidated-bs")
def consolidated_bs(period: Optional[str] = None, user: User = Depends(require_accounting),
                    db: Session = Depends(get_db)):
    """Group balance sheet with intercompany elimination: Intercompany
    Receivable (1180, asset) is eliminated against Intercompany Payable
    (2180, liability) up to the smaller side — group-internal balances are
    not real assets/liabilities of the group."""
    groups: dict[str, dict[str, Decimal]] = {"asset": {}, "liability": {}, "equity": {}}
    names: dict[str, str] = {}
    retained = ZERO
    for line, acc, _e in _all_posted_lines(db, _period_end(period)):
        names.setdefault(acc.code, acc.name)
        if acc.account_type in groups:
            bal = (line.debit or ZERO) - (line.credit or ZERO)
            if acc.account_type != "asset":
                bal = -bal
            groups[acc.account_type][acc.code] = groups[acc.account_type].get(acc.code, ZERO) + bal
        elif acc.account_type == "income":
            retained += (line.credit or ZERO) - (line.debit or ZERO)
        elif acc.account_type == "expense":
            retained -= (line.debit or ZERO) - (line.credit or ZERO)
    ic_asset = groups["asset"].get("1180", ZERO)
    ic_liab = groups["liability"].get("2180", ZERO)
    elimination = min(ic_asset, ic_liab)
    if elimination > 0:
        groups["asset"]["1180"] = ic_asset - elimination
        groups["liability"]["2180"] = ic_liab - elimination

    def section(kind):
        items = [{"code": c, "name": names.get(c, c), "amount": fmt(v)}
                 for c, v in sorted(groups[kind].items()) if v != 0]
        return items, sum(groups[kind].values(), ZERO)

    assets, ta = section("asset")
    liabs, tl = section("liability")
    equity, te = section("equity")
    te += retained
    equity.append({"code": "—", "name": "Retained Earnings (group result)",
                   "amount": fmt(retained)})
    return {
        "assets": assets, "total_assets": fmt(ta),
        "liabilities": liabs, "total_liabilities": fmt(tl),
        "equity": equity, "total_equity": fmt(te),
        "total_liabilities_equity": fmt(tl + te),
        "intercompany_eliminated": fmt(elimination),
        "balanced": ta == tl + te,
    }


@router.get("/reports/bu-comparison")
def bu_comparison(period: Optional[str] = None, user: User = Depends(require_accounting),
                  db: Session = Depends(get_db)):
    """Per-BU P&L for a period — the group scoreboard."""
    per_bu: dict[int, dict[str, Decimal]] = {}
    for line, acc, e in _all_posted_lines(db, _period_end(period)):
        if period and period_of(e.entry_date) != period:
            continue
        row = per_bu.setdefault(e.business_unit_id, {"income": ZERO, "expense": ZERO})
        if acc.account_type == "income":
            row["income"] += (line.credit or ZERO) - (line.debit or ZERO)
        elif acc.account_type == "expense":
            row["expense"] += (line.debit or ZERO) - (line.credit or ZERO)
    names = {bu.id: bu.name for bu in db.query(BusinessUnit).all()}
    items = [{
        "business_unit_id": bu_id, "bu_name": names.get(bu_id, str(bu_id)),
        "income": fmt(v["income"]), "expenses": fmt(v["expense"]),
        "net_profit": fmt(v["income"] - v["expense"]),
    } for bu_id, v in per_bu.items()]
    items.sort(key=lambda x: -x["net_profit"])
    return {"items": items}


# ---------------------------------------------------------------- periods & rules

@router.get("/periods")
def list_periods(business_unit_id: int, user: User = Depends(require_accounting),
                 db: Session = Depends(get_db)):
    rows = db.query(FiscalPeriod).filter(
        FiscalPeriod.business_unit_id == business_unit_id).order_by(FiscalPeriod.period.desc()).all()
    return {"items": [{"period": r.period, "status": r.status,
                       "closed_at": r.closed_at.isoformat() if r.closed_at else None}
                      for r in rows]}


class PeriodAction(BaseModel):
    business_unit_id: int
    period: str = Field(pattern=r"^\d{4}-\d{2}$")


@router.post("/periods/close")
def close_period(payload: PeriodAction, user: User = Depends(require_accounting),
                 db: Session = Depends(get_db)):
    row = db.query(FiscalPeriod).filter(
        FiscalPeriod.business_unit_id == payload.business_unit_id,
        FiscalPeriod.period == payload.period).first()
    if not row:
        row = FiscalPeriod(business_unit_id=payload.business_unit_id, period=payload.period)
        db.add(row)
    row.status = "closed"
    row.closed_by_id = user.id
    row.closed_at = datetime.utcnow()
    db.commit()
    return {"ok": True, "period": payload.period, "status": "closed"}


@router.post("/periods/reopen")
def reopen_period(payload: PeriodAction, user: User = Depends(require_admin),
                  db: Session = Depends(get_db)):
    row = db.query(FiscalPeriod).filter(
        FiscalPeriod.business_unit_id == payload.business_unit_id,
        FiscalPeriod.period == payload.period).first()
    if row:
        row.status = "open"
        row.closed_by_id = None
        row.closed_at = None
        db.commit()
    return {"ok": True, "period": payload.period, "status": "open"}


@router.get("/posting-rules")
def list_posting_rules(business_unit_id: Optional[int] = None,
                       user: User = Depends(require_accounting), db: Session = Depends(get_db)):
    q = db.query(PostingRule)
    if business_unit_id is not None:
        q = q.filter(PostingRule.business_unit_id == business_unit_id)
    rows = q.all()
    return {"items": [{"id": r.id, "business_unit_id": r.business_unit_id,
                       "source_field": r.source_field,
                       "debit_account_code": r.debit_account_code,
                       "credit_account_code": r.credit_account_code} for r in rows],
            "defaults": [{"source_field": f, "debit_account_code": d, "credit_account_code": c}
                         for f, d, c in DEFAULT_POSTING_RULES]}


class PostingRuleIn(BaseModel):
    business_unit_id: Optional[int] = None
    source_field: str
    debit_account_code: str
    credit_account_code: str


@router.put("/posting-rules")
def upsert_posting_rule(payload: PostingRuleIn, user: User = Depends(require_accounting),
                        db: Session = Depends(get_db)):
    row = db.query(PostingRule).filter(
        PostingRule.business_unit_id == payload.business_unit_id,
        PostingRule.source_field == payload.source_field).first()
    if not row:
        row = PostingRule(business_unit_id=payload.business_unit_id,
                          source_field=payload.source_field)
        db.add(row)
    row.debit_account_code = payload.debit_account_code
    row.credit_account_code = payload.credit_account_code
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- opening balances

@router.post("/opening-balance")
async def opening_balance(business_unit_id: int, as_of: date,
                          file: UploadFile = File(...),
                          user: User = Depends(require_accounting),
                          db: Session = Depends(get_db)):
    """Excel upload: columns account_code | debit | credit (header row).

    Creates ONE posted 'opening' journal entry dated `as_of`. The sheet must
    balance overall (Σdebit = Σcredit) like a real opening trial balance.
    """
    from io import BytesIO

    from openpyxl import load_workbook

    bu_or_404(db, business_unit_id)
    ensure_period_open(db, business_unit_id, as_of)
    content = await file.read()
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(400, "Could not read the Excel file")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    # Map columns by header name so both the 3-column layout
    # (account_code|debit|credit) and the template layout
    # (account_code|account_name|debit|credit) work.
    header = [str(c or "").strip().lower() for c in (rows[0] if rows else [])]
    try:
        ci_code = header.index("account_code")
        ci_dr = header.index("debit")
        ci_cr = header.index("credit")
    except ValueError:
        ci_code, ci_dr, ci_cr = 0, 1, 2
    lines = []
    for i, row in enumerate(rows):
        if i == 0:
            continue  # header
        if not row or not row[ci_code]:
            continue
        code = str(row[ci_code]).strip()
        acc = account_by_code(db, business_unit_id, code)
        if not acc:
            raise HTTPException(400, f"Row {i + 1}: account code {code} not found in this BU")
        if acc.is_header:
            raise HTTPException(400, f"Row {i + 1}: {code} is a header account")
        try:
            dr = Decimal(str(row[ci_dr] or 0))
            cr = Decimal(str(row[ci_cr] or 0)) if len(row) > ci_cr else ZERO
        except Exception:
            raise HTTPException(400, f"Row {i + 1}: debit/credit must be numbers")
        if dr == 0 and cr == 0:
            continue
        lines.append({"account_id": acc.id, "debit": dr, "credit": cr})
    if not lines:
        raise HTTPException(400, "No balance rows found")
    td = sum(Decimal(str(l["debit"])) for l in lines)
    tc = sum(Decimal(str(l["credit"])) for l in lines)
    if td != tc:
        raise HTTPException(400, f"Opening balances do not balance: debit {td} != credit {tc}")
    e = create_entry(db, business_unit_id, as_of,
                     f"Opening balances as of {as_of.isoformat()}", "opening",
                     user.id, lines, post=True)
    db.commit()
    db.refresh(e)
    return serialize_entry(db, e)


# ------------------------------------------------------- Phase 6: fixed assets & banking

from src.models import BankAccount, BankTransaction, FixedAsset


def serialize_asset(a: FixedAsset) -> dict:
    book = Decimal(str(a.cost or 0)) - Decimal(str(a.accumulated_depreciation or 0))
    monthly = ZERO
    if a.useful_life_months:
        monthly = (Decimal(str(a.cost or 0)) - Decimal(str(a.salvage_value or 0))) / Decimal(a.useful_life_months)
    return {
        "id": a.id, "business_unit_id": a.business_unit_id, "name": a.name,
        "category": a.category,
        "purchase_date": a.purchase_date.isoformat() if a.purchase_date else None,
        "cost": fmt(a.cost), "salvage_value": fmt(a.salvage_value),
        "useful_life_months": a.useful_life_months,
        "accumulated_depreciation": fmt(a.accumulated_depreciation),
        "book_value": fmt(book), "monthly_depreciation": fmt(monthly),
        "last_depreciated_period": a.last_depreciated_period, "status": a.status,
    }


class FixedAssetIn(BaseModel):
    business_unit_id: int
    name: str = Field(min_length=1, max_length=150)
    category: Optional[str] = None
    purchase_date: date
    cost: float = Field(gt=0)
    salvage_value: float = 0
    useful_life_months: int = Field(gt=0)


@router.get("/fixed-assets")
def list_fixed_assets(business_unit_id: int, user: User = Depends(require_accounting),
                      db: Session = Depends(get_db)):
    rows = db.query(FixedAsset).filter(
        FixedAsset.business_unit_id == business_unit_id).order_by(FixedAsset.purchase_date).all()
    return {"items": [serialize_asset(a) for a in rows]}


@router.post("/fixed-assets")
def create_fixed_asset(payload: FixedAssetIn, user: User = Depends(require_accounting),
                       db: Session = Depends(get_db)):
    bu_or_404(db, payload.business_unit_id)
    a = FixedAsset(**payload.model_dump(), created_by_id=user.id)
    db.add(a)
    db.commit()
    db.refresh(a)
    return serialize_asset(a)


@router.post("/fixed-assets/{asset_id}/dispose")
def dispose_fixed_asset(asset_id: int, user: User = Depends(require_accounting),
                        db: Session = Depends(get_db)):
    a = db.query(FixedAsset).filter(FixedAsset.id == asset_id).first()
    if not a:
        raise HTTPException(404, "Asset not found")
    a.status = "disposed"
    db.commit()
    return serialize_asset(a)


class DepreciateIn(BaseModel):
    business_unit_id: int
    period: str = Field(pattern=r"^\d{4}-\d{2}$")


@router.post("/fixed-assets/depreciate")
def run_depreciation(payload: DepreciateIn, user: User = Depends(require_accounting),
                     db: Session = Depends(get_db)):
    """One month of straight-line depreciation for every active asset that
    has not been depreciated for this period yet. Idempotent per period."""
    assets = db.query(FixedAsset).filter(
        FixedAsset.business_unit_id == payload.business_unit_id,
        FixedAsset.status == "active").all()
    total = ZERO
    count = 0
    for a in assets:
        if a.last_depreciated_period == payload.period:
            continue
        depreciable = Decimal(str(a.cost or 0)) - Decimal(str(a.salvage_value or 0))
        remaining = depreciable - Decimal(str(a.accumulated_depreciation or 0))
        if remaining <= 0:
            continue
        monthly = depreciable / Decimal(a.useful_life_months)
        dep = min(monthly, remaining).quantize(Decimal("0.01"))
        a.accumulated_depreciation = Decimal(str(a.accumulated_depreciation or 0)) + dep
        a.last_depreciated_period = payload.period
        total += dep
        count += 1
    db.commit()
    journal_posted = False
    if total > 0:
        try:
            dep_acc = account_by_code(db, payload.business_unit_id, "6900")
            acc_acc = account_by_code(db, payload.business_unit_id, "1290")
            if dep_acc and acc_acc:
                y, m = payload.period.split("-")
                create_entry(db, payload.business_unit_id, date(int(y), int(m), 1),
                             f"Depreciation {payload.period}", "depreciation", user.id,
                             [{"account_id": dep_acc.id, "debit": total, "credit": 0},
                              {"account_id": acc_acc.id, "debit": 0, "credit": total}],
                             post=True)
                db.commit()
                journal_posted = True
        except Exception:
            db.rollback()
    return {"ok": True, "assets_depreciated": count, "total": fmt(total),
            "journal_posted": journal_posted}


def serialize_bank_account(db: Session, b: BankAccount) -> dict:
    stmt_balance = sum(((Decimal(str(t.deposit or 0)) - Decimal(str(t.withdrawal or 0)))
                        for t in b.transactions), ZERO)
    return {
        "id": b.id, "business_unit_id": b.business_unit_id, "bank_name": b.bank_name,
        "account_name": b.account_name, "account_number": b.account_number,
        "account_id": b.account_id,
        "account_code": b.account.code if b.account else None,
        "statement_balance": fmt(stmt_balance),
        "unreconciled_count": sum(1 for t in b.transactions if not t.is_reconciled),
    }


class BankAccountIn(BaseModel):
    business_unit_id: int
    bank_name: str = Field(min_length=1, max_length=120)
    account_name: Optional[str] = None
    account_number: Optional[str] = None
    account_id: Optional[int] = None


@router.get("/bank-accounts")
def list_bank_accounts(business_unit_id: int, user: User = Depends(require_accounting),
                       db: Session = Depends(get_db)):
    rows = db.query(BankAccount).filter(
        BankAccount.business_unit_id == business_unit_id,
        BankAccount.is_active == True).all()  # noqa: E712
    return {"items": [serialize_bank_account(db, b) for b in rows]}


@router.post("/bank-accounts")
def create_bank_account(payload: BankAccountIn, user: User = Depends(require_accounting),
                        db: Session = Depends(get_db)):
    bu_or_404(db, payload.business_unit_id)
    b = BankAccount(**payload.model_dump(), created_by_id=user.id)
    db.add(b)
    db.commit()
    db.refresh(b)
    return serialize_bank_account(db, b)


class BankTxnIn(BaseModel):
    txn_date: date
    description: Optional[str] = None
    deposit: float = 0
    withdrawal: float = 0


@router.get("/bank-accounts/{bank_id}/transactions")
def list_bank_txns(bank_id: int, user: User = Depends(require_accounting),
                   db: Session = Depends(get_db)):
    b = db.query(BankAccount).filter(BankAccount.id == bank_id).first()
    if not b:
        raise HTTPException(404, "Bank account not found")
    rows = sorted(b.transactions, key=lambda t: (t.txn_date, t.id))
    return {"items": [{"id": t.id, "txn_date": t.txn_date.isoformat(),
                       "description": t.description, "deposit": fmt(t.deposit),
                       "withdrawal": fmt(t.withdrawal),
                       "is_reconciled": t.is_reconciled} for t in rows]}


@router.post("/bank-accounts/{bank_id}/transactions")
def add_bank_txn(bank_id: int, payload: BankTxnIn, user: User = Depends(require_accounting),
                 db: Session = Depends(get_db)):
    b = db.query(BankAccount).filter(BankAccount.id == bank_id).first()
    if not b:
        raise HTTPException(404, "Bank account not found")
    t = BankTransaction(bank_account_id=bank_id, txn_date=payload.txn_date,
                        description=payload.description,
                        deposit=Decimal(str(payload.deposit)),
                        withdrawal=Decimal(str(payload.withdrawal)),
                        created_by_id=user.id)
    db.add(t)
    db.commit()
    db.refresh(t)
    return {"id": t.id, "ok": True}


class ReconcileIn(BaseModel):
    is_reconciled: bool


@router.patch("/bank-transactions/{txn_id}")
def set_reconciled(txn_id: int, payload: ReconcileIn, user: User = Depends(require_accounting),
                   db: Session = Depends(get_db)):
    t = db.query(BankTransaction).filter(BankTransaction.id == txn_id).first()
    if not t:
        raise HTTPException(404, "Transaction not found")
    t.is_reconciled = payload.is_reconciled
    db.commit()
    return {"ok": True}


@router.get("/bank-accounts/{bank_id}/reconciliation")
def bank_reconciliation(bank_id: int, user: User = Depends(require_accounting),
                        db: Session = Depends(get_db)):
    """Statement balance vs the linked ledger account balance."""
    b = db.query(BankAccount).filter(BankAccount.id == bank_id).first()
    if not b:
        raise HTTPException(404, "Bank account not found")
    info = serialize_bank_account(db, b)
    ledger_balance = None
    if b.account_id:
        total = ZERO
        for line, acc, _e in _posted_lines(db, b.business_unit_id):
            if acc.id == b.account_id:
                total += (line.debit or ZERO) - (line.credit or ZERO)
        ledger_balance = fmt(total)
    return {
        **info, "ledger_balance": ledger_balance,
        "difference": (fmt(Decimal(str(ledger_balance)) - Decimal(str(info["statement_balance"])))
                       if ledger_balance is not None else None),
    }


# ------------------------------------------------------- go-live: opening balance template & report exports

@router.get("/opening-balance/template")
def opening_balance_template(business_unit_id: int, user: User = Depends(require_accounting),
                             db: Session = Depends(get_db)):
    """Excel template listing the BU's postable accounts — accountants fill
    in debit/credit opening balances and upload it back."""
    from io import BytesIO

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook

    bu = bu_or_404(db, business_unit_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Opening Balances"
    ws.append(["account_code", "account_name", "debit", "credit"])
    accounts = db.query(Account).filter(
        Account.business_unit_id == business_unit_id,
        Account.is_active == True, Account.is_header == False).order_by(Account.code).all()  # noqa: E712
    for a in accounts:
        ws.append([a.code, a.name, None, None])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="opening_balances_{bu.code or bu.id}.xlsx"'})


def _export_rows(report: str, db: Session, user: User, business_unit_id: Optional[int],
                 period: Optional[str]) -> tuple[str, list[str], list[list]]:
    """(title, header, rows) for the requested report — shared by xlsx + pdf."""
    if report == "tb":
        d = trial_balance(business_unit_id=business_unit_id, period=period, user=user, db=db)
        rows = [[i["code"], i["name"], i["debit"], i["credit"], f'{i["balance"]} {i["balance_side"]}']
                for i in d["items"]]
        rows.append(["", "TOTAL", d["total_debit"], d["total_credit"], ""])
        return "Trial Balance", ["Code", "Account", "Debit", "Credit", "Balance"], rows
    if report == "pl":
        d = profit_loss(business_unit_id=business_unit_id, period=period, user=user, db=db)
        rows = [["INCOME", "", ""]]
        rows += [[i["code"] + " " + i["name"], "", i["amount"]] for i in d["income"]]
        rows.append(["", "Total Income", d["total_income"]])
        rows.append(["EXPENSES", "", ""])
        rows += [[i["code"] + " " + i["name"], "", i["amount"]] for i in d["expenses"]]
        rows.append(["", "Total Expenses", d["total_expenses"]])
        rows.append(["", "NET PROFIT", d["net_profit"]])
        return "Profit & Loss", ["Account", "", "Amount"], rows
    if report in ("bs", "consolidated-bs"):
        d = (balance_sheet(business_unit_id=business_unit_id, period=period, user=user, db=db)
             if report == "bs" else consolidated_bs(period=period, user=user, db=db))
        rows = [["ASSETS", "", ""]]
        rows += [[i["code"] + " " + i["name"], "", i["amount"]] for i in d["assets"]]
        rows.append(["", "Total Assets", d["total_assets"]])
        rows.append(["LIABILITIES", "", ""])
        rows += [[i["code"] + " " + i["name"], "", i["amount"]] for i in d["liabilities"]]
        rows.append(["", "Total Liabilities", d["total_liabilities"]])
        rows.append(["EQUITY", "", ""])
        rows += [[i["code"] + " " + i["name"], "", i["amount"]] for i in d["equity"]]
        rows.append(["", "Total Equity", d["total_equity"]])
        rows.append(["", "Total Liabilities + Equity", d["total_liabilities_equity"]])
        return ("Balance Sheet" if report == "bs" else "Consolidated Balance Sheet"), ["Account", "", "Amount"], rows
    if report == "consolidated-pl":
        d = consolidated_pl(period=period, user=user, db=db)
        rows = [["INCOME", "", ""]]
        rows += [[i["code"] + " " + i["name"], "", i["amount"]] for i in d["income"]]
        rows.append(["", "Total Income", d["total_income"]])
        rows.append(["EXPENSES", "", ""])
        rows += [[i["code"] + " " + i["name"], "", i["amount"]] for i in d["expenses"]]
        rows.append(["", "Total Expenses", d["total_expenses"]])
        rows.append(["", "GROUP NET PROFIT", d["net_profit"]])
        return "Consolidated Profit & Loss", ["Account", "", "Amount"], rows
    if report == "journal":
        q = db.query(JournalEntry).filter(
            JournalEntry.business_unit_id == business_unit_id,
            JournalEntry.status == "posted")
        entries = q.order_by(JournalEntry.entry_date, JournalEntry.id).all()
        if period:
            entries = [e for e in entries if period_of(e.entry_date) == period]
        rows = []
        for e in entries:
            for l in e.lines:
                rows.append([e.entry_date.isoformat(), e.entry_no,
                             l.account.code if l.account else "",
                             l.account.name if l.account else "",
                             e.description or "", fmt(l.debit), fmt(l.credit)])
        return "Journal", ["Date", "Entry No", "Code", "Account", "Description", "Debit", "Credit"], rows
    raise HTTPException(400, f"Unknown report: {report}")


@router.get("/reports/export")
def export_report(report: str, format: str = "xlsx", business_unit_id: Optional[int] = None,
                  period: Optional[str] = None, user: User = Depends(get_current_user),
                  db: Session = Depends(get_db)):
    """Download P&L / balance sheet / trial balance / journal as Excel or PDF."""
    from io import BytesIO

    from fastapi.responses import StreamingResponse

    if report not in ("pl", "bs", "tb", "journal", "consolidated-pl", "consolidated-bs"):
        raise HTTPException(400, "report must be pl, bs, tb, journal, consolidated-pl or consolidated-bs")
    if not report.startswith("consolidated"):
        if not business_unit_id:
            raise HTTPException(400, "business_unit_id is required for this report")
        check_bu_access(user, business_unit_id)
    elif user.role == "CEO":
        raise HTTPException(403, "Consolidated reports are group-level")
    bu = bu_or_404(db, business_unit_id) if business_unit_id else None
    title, header, rows = _export_rows(report, db, user, business_unit_id, period)
    subtitle = f"{bu.name if bu else 'ABC Group'} — {period or 'all periods'}"
    fname_base = f"{report}_{(bu.code or 'group') if bu else 'group'}_{period or 'all'}"

    if format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        ws.append([title])
        ws.append([subtitle])
        ws.append([])
        ws.append(header)
        for cell in ws[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F3864")
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            letter = col[0].column_letter
            ws.column_dimensions[letter].width = max(14, min(46, max((len(str(c.value or "")) for c in col), default=14) + 2))
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{fname_base}.xlsx"'})

    if format == "pdf":
        from fpdf import FPDF

        def safe(t):
            return str(t if t is not None else "").encode("latin-1", errors="replace").decode("latin-1")

        pdf = FPDF(orientation="P", unit="mm", format="A4")
        pdf.add_page()
        pdf.set_font("helvetica", "B", 15)
        pdf.cell(0, 9, safe(title), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("helvetica", "", 10)
        pdf.cell(0, 6, safe(subtitle), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
        n = len(header)
        widths = [190 / n] * n
        if n >= 5:
            widths = [22, 18, 14, 42, 44, 25, 25][:n]
            widths = [w * 190 / sum(widths) for w in widths]
        pdf.set_font("helvetica", "B", 9)
        for i, h in enumerate(header):
            pdf.cell(widths[i], 7, safe(h), border=1)
        pdf.ln()
        pdf.set_font("helvetica", "", 9)
        for r in rows:
            for i, v in enumerate(r):
                txt = f"{v:,.2f}" if isinstance(v, (int, float)) else safe(v)
                pdf.cell(widths[i], 6.5, txt[:44], border=1)
            pdf.ln()
        buf = BytesIO(pdf.output())
        return StreamingResponse(
            buf, media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{fname_base}.pdf"'})

    raise HTTPException(400, "format must be xlsx or pdf")
